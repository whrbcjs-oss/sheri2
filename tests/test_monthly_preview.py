from decimal import Decimal
import unittest
import openpyxl
from openpyxl.worksheet.table import Table
from monthly_report_preview import WorkbookCalculator, REPORT, LEDGER

class MonthlyFormulaTests(unittest.TestCase):
    def setUp(self):
        self.book=openpyxl.Workbook()
        self.report=self.book.active;self.report.title=REPORT
        ledger=self.book.create_sheet(LEDGER)
        for col in range(1,23):ledger.cell(3,col,'column'+str(col))
        for col,label in [(4,'월'),(17,'대분류'),(18,'중분류'),(20,'VAT')]:ledger.cell(3,col,label)
        records=[('2026-01-10','데크거더','공사수입',1100,0,100),
                 ('2026-01-11','데크거더','급여 등',0,550,50),
                 ('2026-01-12','다른부서','급여 등',0,220,20),
                 ('2026-02-10','데크거더','공사수입',2200,0,200)]
        for row,values in enumerate(records,4):
            date,department,category,inflow,outflow,vat=values
            ledger.cell(row,3,date);ledger.cell(row,4,f'=MONTH(C{row})')
            for col,value in [(11,inflow),(12,outflow),(17,department),(18,category),(20,vat)]:ledger.cell(row,col,value)
        # Whole-column SUMIFS must skip an unmatched summary formula, not evaluate it.
        ledger['K1']='=SUBTOTAL(9,K4:K7)'
        ledger.add_table(Table(displayName='표1',ref='A3:V7'))
        self.report['C7']=1;self.report['B11']='급여 등'
        self.calc=WorkbookCalculator(self.book)

    def test_income_minus_vat_and_month_selection(self):
        formula='=SUMIFS(보통예금_원장!$K:$K,보통예금_원장!$Q:$Q,"데크거더",보통예금_원장!$R:$R,"공사수입",보통예금_원장!$D:$D,C7)-SUMIFS(보통예금_원장!$T:$T,보통예금_원장!$Q:$Q,"데크거더",보통예금_원장!$R:$R,"공사수입",보통예금_원장!$D:$D,C7)'
        self.assertEqual(self.calc.expression(formula,REPORT),1000)

    def test_structured_vat_reference_preserves_original_scope(self):
        formula=f'=SUMIFS(보통예금_원장!L:L,보통예금_원장!Q:Q,"데크거더",보통예금_원장!R:R,B11,보통예금_원장!D:D,C7)-SUMIFS(표1[[#All],[VAT]],표1[[#All],[中分類]],B11,표1[[#All],[월]],C7)'
        formula=formula.replace('[中分類]','[중분류]')
        # Includes 20 of VAT from another department, exactly as the original formula would.
        self.assertEqual(self.calc.expression(formula,REPORT),480)

    def test_percentage_text_sum_and_precedence(self):
        self.report['C8']=1000;self.report['C9']=200;self.report['C10']='=SUM(C8:C9)'
        self.report['C11']='text note';self.report['C12']=50
        self.assertEqual(self.calc.expression('=SUM(C8:C12)',REPORT),2450)
        self.assertEqual(self.calc.expression('=C10*4%',REPORT),Decimal('48'))
        self.assertEqual(self.calc.expression('=(10+2)*3-1',REPORT),35)

    def test_explicit_manual_adjustments(self):
        self.assertEqual(self.calc.expression('=100+20+30-10',REPORT),140)

    def test_sheet_name_with_parentheses(self):
        self.report['C10']=42
        self.assertEqual(self.calc.expression(f"='{REPORT}'!$C$10+8",LEDGER),50)

    def test_unsupported_formula_rejected(self):
        with self.assertRaises(ValueError):self.calc.expression('=WEBSERVICE("https://example.com")',REPORT)

    def test_html_payload_cannot_close_script(self):
        from monthly_report_preview import render_html
        html=render_html({'source':'</script><script>alert(1)</script>'})
        self.assertNotIn('"source": "</script>',html)
        self.assertIn('<\\/script>',html)

if __name__=='__main__':unittest.main()
