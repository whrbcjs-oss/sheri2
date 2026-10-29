from pathlib import Path
import hashlib
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
import openpyxl
from monthly_ledger import read_ledger

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'

class LedgerFormulaCacheTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def workbook(self,row=4):
        book=openpyxl.Workbook();ws=book.active;ws.title='보통예금_원장'
        for col,label in [(3,'날짜'),(7,'프로젝트명'),(11,'차변'),(12,'대변'),(17,'대분류'),(18,'중분류'),(20,'VAT')]:ws.cell(3,col,label)
        for col,value in [(3,'2025-01-10'),(7,'sample project'),(11,1100),(12,0),(17,f'=XLOOKUP("sample",G{row}:G{row},Q{row}:Q{row})'),(18,'공사수입'),(20,100)]:ws.cell(row,col,value)
        path=self.root/'bank.xlsx';book.save(path)
        return path

    def cache(self,path,values):
        with zipfile.ZipFile(path) as z:items={n:z.read(n) for n in z.namelist()}
        tree=ET.fromstring(items['xl/worksheets/sheet1.xml'])
        for cell in tree.iter('{'+NS+'}c'):
            if cell.attrib.get('r') not in values:continue
            kind,value=values[cell.attrib['r']]
            cell.set('t',kind)
            node=cell.find('{'+NS+'}v')
            if node is None:node=ET.SubElement(cell,'{'+NS+'}v')
            node.text=str(value) if value is not None else None
        items['xl/worksheets/sheet1.xml']=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in items.items():z.writestr(name,data)

    def test_unsupported_formula_at_row_8569_uses_saved_excel_result(self):
        path=self.workbook(8569);self.cache(path,{'Q8569':('str','데크거더')})
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        data=read_ledger(path)
        self.assertEqual(len(data['records']),1)
        self.assertEqual(data['records'][0]['department'],'데크거더')
        self.assertEqual(data['records'][0]['inflow'],'1100')
        self.assertEqual(data['cached_formula_count'],1)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)

    def test_missing_cached_result_reports_cell_and_recalculation_steps(self):
        path=self.workbook()
        with self.assertRaisesRegex(ValueError,'Q4.*XLOOKUP.*Ctrl\\+Alt\\+F9'):read_ledger(path)

    def test_zero_and_empty_string_cached_results_are_valid(self):
        path=self.workbook();book=openpyxl.load_workbook(path);ws=book.active
        ws['T4']='=XLOOKUP("missing",G4:G4,T4:T4)'
        ws['R4']='=XLOOKUP("sample",G4:G4,R4:R4)';book.save(path)
        self.cache(path,{'Q4':('str','데크거더'),'T4':('n',0),'R4':('str',None)})
        data=read_ledger(path)
        self.assertEqual(data['records'][0]['vat'],'0')
        self.assertEqual(data['records'][0]['category'],'')
        self.assertEqual(data['cached_formula_count'],3)

    def test_excel_error_is_not_silently_treated_as_zero(self):
        path=self.workbook();self.cache(path,{'Q4':('e','#VALUE!')})
        with self.assertRaisesRegex(ValueError,'Q4.*#VALUE!'):read_ledger(path)

    def test_supported_parent_recalculates_using_cached_child(self):
        path=self.workbook();book=openpyxl.load_workbook(path);ws=book.active
        ws['T4']='=XLOOKUP("sample",G4:G4,T4:T4)';ws['K4']='=T4*11';book.save(path)
        self.cache(path,{'Q4':('str','데크거더'),'T4':('n',100),'K4':('n',9999)})
        data=read_ledger(path)
        # Supported arithmetic is recomputed instead of trusting its stale cached result.
        self.assertEqual(data['records'][0]['inflow'],'1100')
        self.assertEqual(data['cached_formula_count'],2)

class SearchClassificationTests(unittest.TestCase):
    setUp=LedgerFormulaCacheTests.setUp
    workbook=LedgerFormulaCacheTests.workbook
    cache=LedgerFormulaCacheTests.cache

    def native_book(self,text,row=8569):
        path=self.workbook(row);book=openpyxl.load_workbook(path);ws=book.active
        ws.cell(row,14,text)
        ws.cell(row,17,f'=IF(OR(ISNUMBER(SEARCH("LB", N{row})), ISNUMBER(SEARCH("김제공장", N{row})), ISNUMBER(SEARCH("기업부설", N{row}))), "데크거더", "0")')
        book.save(path)
        return path

    def test_exact_user_formula_without_cache(self):
        for text in ['LB 공사','김제공장 운영','기업부설 연구소','lb 소문자']:
            with self.subTest(text=text):
                result=read_ledger(self.native_book(text))
                self.assertEqual(result['records'][0]['department'],'데크거더')
                self.assertEqual(result['native_formula_count'],1)
                self.assertEqual(result['cached_formula_count'],0)

    def test_no_search_match_isnumber_is_false(self):
        result=read_ledger(self.native_book('기타 현장'))
        self.assertEqual(result['records'][0]['department'],'0')

    def test_native_classification_recomputes_stale_cache(self):
        path=self.native_book('LB project');self.cache(path,{'Q8569':('str','0')})
        self.assertEqual(read_ledger(path)['records'][0]['department'],'데크거더')

    def test_if_does_not_evaluate_unused_error_branch(self):
        path=self.workbook();book=openpyxl.load_workbook(path)
        book.active['Q4']='=IF(TRUE,"데크거더",SEARCH("missing","text"))';book.save(path)
        self.assertEqual(read_ledger(path)['records'][0]['department'],'데크거더')

if __name__=='__main__':unittest.main()
