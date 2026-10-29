from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
import openpyxl
from monthly_ledger import money,read_ledger

class MoneyParsingTests(unittest.TestCase):
    def test_blank_and_whitespace_values_are_zero(self):
        for value in [None,'',' ','  ','\t\r\n','\xa0','\u202f','\u3000','\u200b','\ufeff','\u2060',' \u200b \ufeff ','\u200e','\u200f','\u202a','\u2066','\u200c','\u200d','\xad']:
            with self.subTest(value=repr(value)):self.assertEqual(money(value),Decimal(0))

    def test_numeric_values_keep_sign_and_fraction(self):
        for value,expected in [('1,234.50','1234.50'),(' (1,234.50) ','-1234.50'),('-1,234','-1234'),('\xa01234\ufeff','1234'),(0,'0')]:
            with self.subTest(value=value):self.assertEqual(money(value),Decimal(expected))

    def test_invalid_text_is_not_silently_zeroed(self):
        for value in ['unknown','-','—',',','1 234','NaN','Infinity','1\u200b234']:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):money(value)

    def fixture(self,path,values):
        book=openpyxl.Workbook();ws=book.active;ws.title='보통예금_원장'
        for column,label in [(3,'날짜'),(11,'차변'),(12,'대변'),(17,'대분류'),(18,'중분류'),(20,'VAT')]:ws.cell(3,column,label)
        row=11822
        for column,value in [(3,'2025-02-10'),(11,values[0]),(12,values[1]),(17,'데크거더'),(18,'공사수입'),(20,values[2])]:ws.cell(row,column,value)
        book.save(path)

    def test_row_11822_with_invisible_blanks_imports(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'bank.xlsx';self.fixture(path,['1100','\xa0',' \u200b '])
            result=read_ledger(path)
            self.assertEqual(len(result['records']),1)
            row=result['records'][0]
            self.assertEqual((row['inflow'],row['outflow'],row['vat']),('1100','0','0'))

    def test_user_outflow_row_keeps_credit_when_debit_and_vat_blank(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'bank.xlsx';self.fixture(path,[' \u200e','2500','\xa0\u200b'])
            row=read_ledger(path)['records'][0]
            self.assertEqual((row['inflow'],row['outflow'],row['vat']),('0','2500','0'))

    def test_invalid_amount_identifies_exact_cell_and_contents(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'bank.xlsx';self.fixture(path,['1100','bad amount','0'])
            with self.assertRaisesRegex(ValueError,"11822행: 보통예금_원장!L11822 \\(대변\\):.*'bad amount'"):read_ledger(path)

    def test_formula_returning_whitespace_imports_as_blank(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'bank.xlsx';self.fixture(path,['1100','=IF(TRUE," ","100")','100'])
            row=read_ledger(path)['records'][0]
            self.assertEqual(row['outflow'],'0')
            self.assertEqual(row['vat'],'100')

if __name__=='__main__':unittest.main()
