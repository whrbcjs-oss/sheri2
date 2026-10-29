from pathlib import Path
import hashlib
import tempfile
import unittest
from unittest.mock import patch
import secrets
import openpyxl
import collaboration as s
import monthly_store as store
from monthly_ledger import read_ledger

class MonthlyStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)
        self.data_patch=patch.object(s,'DATA',self.root/'data');self.data_patch.start();self.addCleanup(self.data_patch.stop)
        self.state_patch=patch.object(s,'STATE',self.root/'private');self.state_patch.start();self.addCleanup(self.state_patch.stop)
        password=secrets.token_urlsafe(20)
        s.create_user(None,'admin','관리자',password,bootstrap=True);self.token=s.login('admin',password)
        s.create_user(self.token,'staff','담당자',password);self.staff=s.login('staff',password)
        self.context={'token':self.token,'user':s.validate_session(self.token),'snapshots':{},'connections':[]}
        reset=s.RUN.set(self.context);self.addCleanup(s.RUN.reset,reset)
        self.addCleanup(self.close)
        self.db=store.connect(2026)

    def close(self):
        for c in self.context['connections']:
            try:c.rollback();c.close()
            except Exception:pass

    def workbook(self,income='1100'):
        book=openpyxl.Workbook();ws=book.active;ws.title='보통예금_원장'
        columns={2:'법인',3:'날짜',4:'월',7:'프로젝트명',9:'거래처명',11:'차변',12:'대변',17:'대분류',18:'중분류',19:'품의내역명',20:'VAT',22:'비고'}
        for col,label in columns.items():ws.cell(3,col,label)
        for r,(date,category,debit,credit,vat,dept) in enumerate([
            ('2026-01-10','공사수입',income,'0','100','데크거더'),
            ('2026-01-11','미불','0','550','50','데크거더'),
            ('2026-02-10','공사수입','2200','0','200','데크거더'),
            ('2026-02-11','미불','0','1100','100','데크거더'),
            ('2026-01-12','미불','0','220','20','다른부서')],4):
            for c,v in [(2,'테스트'),(3,date),(4,'=MONTH(C'+str(r)+')'),(11,debit),(12,credit),(17,dept),(18,category),(20,vat)]:ws.cell(r,c,v)
        path=self.root/f'input-{income}.xlsx';book.save(path)
        return path

    def save(self,payload,blob,month):return store.save_month(self.db,2026,month,payload,blob,self.token)

    def test_ledger_sheet_only_and_month_append_replace(self):
        path=self.workbook();payload=read_ledger(path);blob=path.read_bytes()
        self.assertEqual(len(payload['records']),5)
        self.save(payload,blob,1);self.save(payload,blob,2)
        feb_before=self.db.execute('SELECT * FROM monthly_ledger WHERE month=2').fetchall()
        path=self.workbook('3300');replacement=read_ledger(path)
        self.save(replacement,path.read_bytes(),1)
        self.assertEqual(self.db.execute('SELECT * FROM monthly_ledger WHERE month=2').fetchall(),feb_before)
        report=store.report(self.db,2026);rows={r['row']:r for r in report['rows']}
        self.assertEqual(rows[8]['values'][0],3200)
        self.assertEqual(rows[8]['values'][1],2000)
        # The other department's VAT must never reduce the deck expense.
        self.assertEqual(rows[17]['values'][0],500)
        self.assertEqual(report['months'],[1,2])
        self.assertIsNone(report['closing'])
        self.assertFalse(report['opening_confirmed'])

    def test_duplicate_noop_and_permissions(self):
        path=self.workbook();payload=read_ledger(path);blob=path.read_bytes()
        self.assertTrue(self.save(payload,blob,1));before=s.revision(self.db)
        self.assertFalse(self.save(payload,blob,1));self.assertEqual(s.revision(self.db),before)
        with self.assertRaises(PermissionError):store.save_month(self.db,2026,2,payload,blob,self.staff)

    def test_settings_and_paired_adjustments(self):
        path=self.workbook();payload=read_ledger(path);self.save(payload,path.read_bytes(),1)
        store.save_settings(self.db,2026,'5000','4')
        store.save_insurance(self.db,2026,1,'100','보험료 보전 시험')
        report=store.report(self.db,2026);rows={r['row']:r for r in report['rows']}
        self.assertEqual(rows[8]['values'][0],1100)
        self.assertEqual(rows[16]['values'][0],100)
        self.assertEqual(rows[25]['values'][0],44)
        self.assertEqual(report['closing'],5456)
        self.assertTrue(self.db.execute('SELECT COUNT(*) FROM change_history').fetchone()[0]>0)
        # Re-uploading a month must not erase the manual inputs.
        replacement=self.workbook('2200');self.save(read_ledger(replacement),replacement.read_bytes(),1)
        self.assertEqual(store.settings(self.db)['opening_amount'],'5000')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM monthly_adjustments').fetchone()[0],2)

    def test_same_month_stale_writer_rejected(self):
        other=store.connect(2026);path=self.workbook();payload=read_ledger(path)
        self.save(payload,path.read_bytes(),1)
        with self.assertRaises(s.StaleData):store.save_month(other,2026,2,payload,path.read_bytes(),self.token)

    def test_backup_restore_separate_database(self):
        finance=s.DATA/'finance_2026.db';finance.write_bytes(b'untouched finance fixture')
        path=self.workbook();payload=read_ledger(path);self.save(payload,path.read_bytes(),1)
        saved=s.backup_database(self.db.path)
        self.save(payload,path.read_bytes(),2)
        s.restore_database(self.token,self.db.path,saved.read_bytes(),s.revision(self.db))
        self.assertEqual([i['month'] for i in store.imports(self.db)],[1])
        self.assertEqual(finance.read_bytes(),b'untouched finance fixture')

    def test_ui_uses_existing_monthly_menu(self):
        from streamlit.testing.v1 import AppTest
        path=self.workbook();payload=read_ledger(path);self.save(payload,path.read_bytes(),1)
        root=Path(__file__).resolve().parents[1]
        at=AppTest.from_file(str(root/'multiuser_app.py'),default_timeout=30)
        at.session_state['_auth_token']=self.token
        at.session_state['main_system']='엘비데크 월별보고'
        at.session_state['lb_monthly_page']='통장입출금기준 현황'
        at.run()
        self.assertFalse(at.exception,str(at.exception))
        self.assertTrue(any('보고서 HTML' in b.label for b in at.get('download_button')))
        at.session_state['lb_monthly_page']='원장 업로드';at.run()
        self.assertFalse(at.exception,str(at.exception))
        # Simulate the parsed upload; the real parser is exercised above.
        at.session_state['_lb_upload']={'digest':payload['sha256'],'payload':payload,'blob':path.read_bytes()}
        at.run()
        next(x for x in at.selectbox if x.label=='저장할 월').set_value(2).run()
        next(x for x in at.checkbox if '원장만 저장' in x.label).check().run()
        next(x for x in at.button if x.label=='선택 월 원장 저장').click().run()
        self.assertFalse(at.exception,str(at.exception))
        self.assertEqual([i['month'] for i in store.imports(self.db)],[1,2])

    def test_staff_can_edit_settings_with_own_audit_identity(self):
        ctx={'token':self.staff,'user':s.validate_session(self.staff),'snapshots':{},'connections':[]}
        reset=s.RUN.set(ctx)
        try:
            db=store.connect(2026)
            store.save_settings(db,2026,'100','4')
            self.assertEqual(db.execute('SELECT actor FROM change_history ORDER BY id DESC LIMIT 1').fetchone()[0],'staff')
        finally:
            for db in ctx['connections']:db.close()
            s.RUN.reset(reset)

    def test_bad_date_and_missing_vat_rejected(self):
        path=self.workbook();book=openpyxl.load_workbook(path);ws=book.active
        ws['C4']='bad date';book.save(path)
        with self.assertRaises(ValueError):read_ledger(path)
        ws['C4']='2026-01-10';ws['T3']='missing';book.save(path)
        with self.assertRaises(ValueError):read_ledger(path)

if __name__=='__main__':unittest.main()
