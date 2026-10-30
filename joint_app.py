"""Joint venture management integrated into the authenticated Streamlit application."""
from datetime import date
from io import BytesIO
import json
import sqlite3
from zipfile import BadZipFile
import pandas as pd
import openpyxl
import streamlit as st
import streamlit.components.v1 as components
import collaboration as shared
import joint_store as db
import joint_workflow as flow
from joint_report import render as report_html,money,period_label
from joint_import import read as import_rows,LABELS

VERSION='2026.10.30.3'

def saved():
    st.session_state['_jv_success']='저장했습니다.';st.rerun()

def month_fields(default=None):
    default=default or date.today().strftime('%Y-%m')
    a,b=st.columns(2)
    return a.text_input('시작 연월 (YYYY-MM)',default),b.text_input('종료 연월 (YYYY-MM)',default)

COMPANY_TYPES=['당사','관계사','외부 회원사','미분류']

def company_page(c):
    st.header('법인·담당자 등록')
    cos=db.companies(c);filters=st.columns([1,2]);filter_=filters[0].selectbox('회사 구분별 조회',['전체']+COMPANY_TYPES)
    visible=[x for x in cos if filter_=='전체' or json.loads(x['info']).get('회사구분','미분류')==filter_]
    selected=filters[1].selectbox('등록·수정할 회사',[None]+[x['id'] for x in visible],format_func=lambda id:'새 회사' if id is None else next(x['name']+' ['+x['code']+']' for x in visible if x['id']==id))
    old=next((x for x in cos if x['id']==selected),{});info=json.loads(old.get('info','{}'))
    with st.form('jv_company_'+str(selected)):
        columns=st.columns(3)
        with columns[0]:
            st.subheader('기본정보')
            pair=st.columns(2);code=pair[0].text_input('법인코드',old.get('code',''));name=pair[1].text_input('법인명',old.get('name',''))
            pair=st.columns(2);kind=pair[0].selectbox('회사 구분',COMPANY_TYPES,index=COMPANY_TYPES.index(info.get('회사구분','미분류')) if info.get('회사구분','미분류') in COMPANY_TYPES else 3)
            values={'회사구분':kind,'사업자등록번호':pair[1].text_input('사업자등록번호',info.get('사업자등록번호','')),'대표자':st.text_input('대표자명',info.get('대표자',''))}
        with columns[1]:
            st.subheader('연락처·계좌')
            pair=st.columns(2);values['전화']=pair[0].text_input('전화번호',info.get('전화',''));values['팩스']=pair[1].text_input('팩스번호',info.get('팩스',''))
            values['주소']=st.text_input('주소',info.get('주소',''));values['입금계좌']=st.text_input('입금계좌',info.get('입금계좌',''))
        with columns[2]:
            st.subheader('담당자')
            for suffix,label in [('', '1'),('2','2')]:
                pair=st.columns(2);values['담당자'+suffix]=pair[0].text_input('담당자 '+label+' 성명',info.get('담당자'+suffix,''));values['직급'+suffix]=pair[1].text_input('담당자 '+label+' 직급',info.get('직급'+suffix,''))
                key='담당자 연락처' if not suffix else '담당자2 연락처';values[key]=st.text_input('담당자 '+label+' 연락처',info.get(key,''))
        if st.form_submit_button('법인 저장',type='primary'):db.company(c,code,name,{**info,**values},selected);saved()
    st.dataframe([{'법인코드':x['code'],'법인명':x['name'],'회사 구분':json.loads(x['info']).get('회사구분','미분류')} for x in visible],hide_index=True,use_container_width=True)

def project_page(c,current=None):
    st.header('프로젝트·회원사 등록')
    cos=db.companies(c)
    if not cos:st.info('먼저 법인·담당자에서 당사와 참여 회사를 등록하세요.');return
    labels={x['id']:f"{x['name']} [{x['code']}] · {json.loads(x['info']).get('회사구분','미분류')}" for x in cos};reverse={v:k for k,v in labels.items()}
    ps=db.projects(c);selected=st.selectbox('프로젝트 등록·수정',[None]+[x['id'] for x in ps],index=([None]+[x['id'] for x in ps]).index(current) if current in [x['id'] for x in ps] else 0,format_func=lambda id:'새 프로젝트' if id is None else next(x['name'] for x in ps if x['id']==id))
    old=next((x for x in ps if x['id']==selected),{});ms=db.members(c,selected) if selected else []
    owner_options=[x['id'] for x in cos if json.loads(x['info']).get('회사구분','미분류')!='외부 회원사' or x['id']==old.get('owner_id')]
    if not owner_options:st.info('관리법인으로 사용할 당사 또는 관계사를 등록하세요.');return
    with st.form('jv_project_'+str(selected)):
        left,right=st.columns(2)
        with left:
            name=st.text_input('공사명',old.get('name',''));owner=st.selectbox('당사·관계사 (관리법인)',owner_options,format_func=labels.get,index=owner_options.index(old.get('owner_id',owner_options[0])))
            contract=st.text_input('계약금액 전체',old.get('contract','0'));supply=st.text_input('공급가액 전체',old.get('supply','0'));vat=st.text_input('계약 부가세액',old.get('vat','0'))
        with right:
            start=st.date_input('공사 시작',date.fromisoformat(old.get('start',date.today().isoformat())));end=st.date_input('공사 종료',date.fromisoformat(old.get('end',date.today().isoformat())))
            st.caption('기존 참여 회사의 연결은 유지합니다. 연결 자료가 있으면 지분 변경·회원사 삭제만 제한되며 공사명·계약정보는 수정할 수 있습니다.')
        initial=[{'법인':labels[m['company_id']],'지분율(%)':m['share'],'관리비 적용':bool(m['fee_enabled']),'관리비율(%)':m['fee_rate']} for m in ms] or [{'법인':labels[owner],'지분율(%)':'100','관리비 적용':False,'관리비율(%)':'0'}]
        st.caption('당사를 포함해 지분 합계 100%로 등록하세요. 법인명 옆에 회사 구분이 표시됩니다.')
        frame=st.data_editor(pd.DataFrame(initial),num_rows='dynamic',hide_index=True,use_container_width=True,column_config={'법인':st.column_config.SelectboxColumn(options=list(reverse),required=True)})
        if st.form_submit_button('프로젝트 저장',type='primary'):
            entries=[dict(company_id=reverse[x['법인']],share=x['지분율(%)'],fee_enabled=x['관리비 적용'],fee_rate=x['관리비율(%)']) for x in frame.to_dict('records') if x.get('법인') in reverse]
            db.project(c,name,owner,contract,supply,vat,start,end,entries,selected);saved()

def ledger_page(c,p,ms):
    st.header('원장 업로드·수기입력')
    tab1,tab2,tab3=st.tabs(['회원사 원가 수기입력','엑셀 업로드','원장 조회'])
    with tab1:
        with st.form('jv_manual_'+str(p)):
            cols=st.columns(3)
            with cols[0]:
                who=st.selectbox('원가투입 회사',[m['company_id'] for m in ms],format_func=lambda id:next(m['name'] for m in ms if m['company_id']==id))
                dt=st.date_input('기표일자');month=st.text_input('배분연월 (YYYY-MM)',date.today().strftime('%Y-%m'))
            with cols[1]:
                rec=dict(posting_date=dt,allocation_month=month,account=st.text_input('계정과목'),vendor=st.text_input('거래처명'),description=st.text_input('적요'),note=st.text_input('비고·정정 사유'))
            with cols[2]:
                for key,label in [('tax_supply','세금계산서 공급가액 (VAT 제외)'),('vat','부가세 (VAT)'),('invoice','계산서 금액 (면세)'),('receipt','기타영수증 금액')]:rec[key]=st.text_input(label,'0')
            if st.form_submit_button('합계 확인'):
                st.info('합계금액: '+money(sum((db.num(rec[k]) for k in ['tax_supply','vat','invoice','receipt']),db.num(0)))+'원')
            if st.form_submit_button('원가 저장',type='primary'):db.add_ledger(c,p,who,[rec]);saved()
    with tab2:
        st.info('아마란스 실제 양식 확인 전에는 열을 직접 지정합니다. 전표번호와 전표행 식별값으로 중복을 차단하며, 중복이 있으면 파일 전체를 저장하지 않습니다.')
        wb=openpyxl.Workbook();ws=wb.active;ws.title='원장';ws.append(list(LABELS.values()))
        stream=BytesIO();wb.save(stream)
        st.download_button('입력용 빈 엑셀 양식',stream.getvalue(),file_name='공동도급_원장입력양식.xlsx')
        up=st.file_uploader('원장 엑셀 선택',type=['xlsx'],key='jv_upload_'+str(p))
        if up:
            blob=up.getvalue();book=openpyxl.load_workbook(BytesIO(blob),read_only=True,data_only=True)
            try:
                sheet=st.selectbox('시트',book.sheetnames);header=st.number_input('제목 행 번호',min_value=1,value=1,step=1)
                last=st.number_input('마지막 자료 행 (0이면 끝까지)',min_value=0,value=0,step=1)
                if last and last<=header:raise ValueError('마지막 자료 행은 제목 행 다음이어야 합니다.')
                cells=next(book[sheet].iter_rows(min_row=int(header),max_row=int(header)),[])
                options={i+1:f'{openpyxl.utils.get_column_letter(i+1)} · {cell.value or "빈 제목"}' for i,cell in enumerate(cells)}
                if not options:st.warning('해당 행에 열이 없습니다.');return
                who=st.selectbox('업로드 원가투입 법인',[m['company_id'] for m in ms],format_func=lambda x:next(m['name'] for m in ms if m['company_id']==x),key='jv_import_company_'+str(p))
                mapping={}
                with st.expander('엑셀 열 연결',expanded=True):
                    for key,label in LABELS.items():
                        match=next((col for col,text in options.items() if text.split(' · ',1)[1]==label),None)
                        choices=[None]+list(options);mapping[key]=st.selectbox(label,choices,index=choices.index(match),format_func=lambda x:'사용 안 함' if x is None else options[x],key='jv_map_'+str(p)+'_'+key)
                if st.button('분석 미리보기'):
                    records=import_rows(blob,sheet,header,mapping,int(last) or None);st.dataframe(pd.DataFrame(records[:100]).rename(columns=LABELS),hide_index=True);st.success(f'{len(records):,}행 분석 완료. 저장하려면 아래 버튼을 누르세요.')
                if st.button('검증 후 원장 전체 저장',type='primary'):
                    records=import_rows(blob,sheet,header,mapping,int(last) or None);db.add_ledger(c,p,who,records,blob,up.name);saved()
            finally:book.close()
    with tab3:
        frame=pd.DataFrame(db.ledger(c,p));labels={'id':'원장번호','company':'원가투입 회사','posting_date':'기표일자','allocation_month':'배분연월','account':'계정과목','vendor':'거래처명','description':'적요','tax_supply':'세금계산서 공급가액','vat':'VAT','invoice':'계산서 금액','receipt':'기타영수증 금액','note':'비고'}
        frame=frame[[k for k in labels if k in frame.columns]].rename(columns=labels);st.dataframe(frame,hide_index=True,use_container_width=True)
        st.download_button('원장 CSV 내려받기',safe_csv(frame),file_name='공동도급_원장.csv')

def csv_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        try:db.num(value)
        except ValueError:return "'"+value
    return value


def safe_csv(frame):
    # Prevent spreadsheet formula execution from text fields in exported CSV.
    frame=frame.copy()
    for col in frame.columns:
        frame[col]=frame[col].map(csv_cell)
    return frame.to_csv(index=False).encode('utf-8-sig')

def allocation_page(c,p,start,end):
    st.header('원가배분내역서')
    st.caption('원장별 공급가액·VAT를 지분 배분합니다. 원 단위 반올림 후 잔여액은 지분이 가장 큰 회사에 반영하며 총액은 보존합니다.')
    if st.button('미배분 원장 자동 배분',type='primary'):
        count=db.allocate(c,p);st.success(f'{count}개 원장을 배분했습니다. 기존 배분·수정액은 유지했습니다.')
    al=db.allocations(c,p,start,end);frame=pd.DataFrame(al)
    labels={'ledger_id':'원장번호','allocation_month':'배분연월','name':'회원사','account':'계정과목','description':'적요','share':'적용 지분율(%)','original_base':'자동계산 공급가액','original_vat':'자동계산 VAT','base':'배분 공급가액','vat':'배분 VAT','reason':'수정 사유'}
    frame=frame[[k for k in labels if k in frame.columns]].rename(columns=labels);st.dataframe(frame,hide_index=True,use_container_width=True)
    st.download_button('원가배분 CSV',safe_csv(frame),file_name='원가배분내역서.csv')
    if al:
        line=st.selectbox('수정할 원장ID',list(dict.fromkeys(a['ledger_id'] for a in al)))
        subset=[{'id':a['id'],'회원사':a['name'],'base':a['base'],'vat':a['vat'],'reason':a['reason']} for a in al if a['ledger_id']==line]
        with st.form('jv_alloc_edit_'+str(line)):
            st.caption('base는 VAT 제외 원가입니다. 모든 회원사 합계가 원장과 각각 일치해야 합니다. 수정할 회원사뿐 아니라 전체 행의 사유를 입력하세요.')
            edits=st.data_editor(pd.DataFrame(subset),disabled=['id','회원사'],hide_index=True,use_container_width=True,column_config={'id':'배분번호','base':'배분 공급가액','vat':'배분 VAT','reason':'수정 사유'})
            if st.form_submit_button('배분 수정 저장'):db.revise(c,p,line,edits.to_dict('records'));saved()

def pick_table(title,records,selected,key,used=None):
    st.subheader(title)
    if not records:st.info('등록된 내역이 없습니다.');return []
    used=used or {};initial=[]
    for r in records:initial.append({'선택':r['id'] in selected,**r,'다른 정산 사용':used.get(r['id'],'')})
    import hashlib
    fingerprint=hashlib.sha256(json.dumps(initial,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:12]
    frame=st.data_editor(pd.DataFrame(initial),hide_index=True,use_container_width=True,disabled=[x for x in initial[0] if x!='선택'],column_config={'id':st.column_config.NumberColumn('내역번호')},key=key+'_'+fingerprint)
    return [int(x['id']) for x in frame.to_dict('records') if x['선택']]

def entry_form(c,p,m,kind,default_month=None):
    es=flow.source_entries(c,p,m['id'],kind);scope=f"{p}_{m['id']}_{kind}"
    if st.session_state.get('entry_choose_'+scope) not in [None]+[e['id'] for e in es]:st.session_state['entry_choose_'+scope]=None
    selected=st.selectbox('등록·수정할 내역',[None]+[e['id'] for e in es],format_func=lambda id:'새 내역' if id is None else str(id)+' · '+next(e['entry_date']+' '+e['data'].get('label',e['data'].get('vendor','')) for e in es if e['id']==id),key='entry_choose_'+scope)
    old=next((e for e in es if e['id']==selected),{});v=old.get('data',{})
    with st.form('entry_form_'+scope+'_'+str(selected)):
        cols=st.columns(3)
        with cols[0]:
            dt=st.date_input({'receipt':'공사대금 수령일','subcontract':'청구일','recovery':'회수일','opening':'입력일'}[kind],date.fromisoformat(old.get('entry_date',date.today().isoformat())))
            month=st.text_input('배분연월 (YYYY-MM)',old.get('period_start',default_month or date.today().strftime('%Y-%m')))
            payload={}
            if kind=='receipt':payload['label']=st.text_input('기성 구분 (분할 수령은 같은 구분)',v.get('label',''))
            if kind=='subcontract':payload['vendor']=st.text_input('하도급 업체명',v.get('vendor',''))
        with cols[1]:
            if kind=='opening':
                payload['arrears']=st.text_input('전월 미수액',v.get('arrears','0'));payload['interest']=st.text_input('지연이자',v.get('interest','0'))
            else:
                payload['supply']=st.text_input('수령 공급가액' if kind=='receipt' else '청구 공급가액' if kind=='subcontract' else '회수 공급가액',v.get('supply','0'))
                payload['vat']=st.text_input('VAT',v.get('vat','0'))
            payload['note']=st.text_area('비고·수정 사유',v.get('note',''),height=90)
        with cols[2]:
            if kind=='subcontract':
                for k,label in [('paid_date','지급일 (YYYY-MM-DD)'),('paid_supply','지급 공급가액'),('paid_vat','지급 VAT'),('work_type','공종명'),('method','지급방식')]:payload[k]=st.text_input(label,v.get(k,'0' if k in ['paid_supply','paid_vat'] else ''))
                if payload['paid_date']:db.day(payload['paid_date'])
            elif kind=='receipt':st.caption('수령 내역을 건별로 등록하세요. 관리비·보험료는 다음 탭에서 수령 범위를 선택해 별도로 정산합니다.')
            elif kind=='recovery':st.caption('하도급 지급액은 자동 반영됩니다. 동일 지급을 회수내역에 중복 입력하지 마세요.')
        if st.form_submit_button('입력 저장',type='primary'):db.save_entry(c,p,m['id'],kind,dt,month,month,payload,selected);saved()
    display=[]
    for e in es:
        v=e['data'];display.append({'id':e['id'],'일자':e['entry_date'],'배분연월':e['period_start'],'구분·업체':v.get('label',v.get('vendor','')),'공급가액':v.get('supply',''),'VAT':v.get('vat',''),'비고':v.get('note','')})
    st.dataframe(display,hide_index=True,use_container_width=True)
    if selected and st.checkbox('선택한 내역을 삭제합니다.',key='delete_check_'+scope):
        if st.button('선택 내역 삭제',key='delete_entry_'+scope):flow.delete_entry(c,p,m['id'],selected);saved()

def fee_page(c,p,m):
    gs=flow.groups(c,p,m['id']);scope=f"{p}_{m['id']}"
    if st.session_state.get('_jv_next_fee',[])[:2]==[p,m['id']]:
        st.session_state['fee_choice_'+scope]=st.session_state.pop('_jv_next_fee')[2]
    if st.session_state.get('fee_choice_'+scope) not in [None]+[g['id'] for g in gs]:st.session_state['fee_choice_'+scope]=None
    selected=st.selectbox('관리비·보험료 정산',[None]+[g['id'] for g in gs],format_func=lambda id:'새 정산' if id is None else next(g['label'] for g in gs if g['id']==id),key='fee_choice_'+scope)
    old=next((g for g in gs if g['id']==selected),{});defaults=json.loads(old.get('receipt_ids','[]'));settings={v['label']:v for v in json.loads(old.get('calculations','[]'))}
    receipts=flow.source_entries(c,p,m['id'],'receipt')
    records=[{'id':e['id'],'수령일':e['entry_date'],'기성 구분':e['data']['label'],'공급가액':e['data']['supply'],'VAT':e['data']['vat']} for e in receipts]
    linked={rid:g['label'] for g in gs if g['id']!=selected for rid in json.loads(g['receipt_ids'])}
    ids=pick_table('정산할 공사대금 수령 범위',records,defaults,'fee_receipts_'+scope+'_'+str(selected),linked)
    stages=list(dict.fromkeys(e['data']['label'] for e in receipts if e['id'] in ids))
    with st.form('fee_form_'+scope+'_'+str(selected)):
        name=st.text_input('관리비·보험료 정산명',old.get('label',''));inputs={}
        tabs=st.tabs(stages) if stages else []
        for stage,tab in zip(stages,tabs):
            with tab:
                values=settings.get(stage,{})
                supply=sum((db.num(e['data']['supply']) for e in receipts if e['id'] in ids and e['data']['label']==stage),db.num(0))
                st.info(stage+' 선택 공급가액 합계: '+money(supply)+'원')
                cols=st.columns(3)
                values_out={'fee_rate':cols[0].text_input('관리비율 (%)',values.get('fee_rate',m['fee_rate']),key='rate_'+scope+'_'+str(selected)+'_'+stage)}
                for i,label in enumerate(db.EXCLUSIONS):values_out[label]=cols[(i+1)%3].text_input(label,values.get(label,'0'),key='fee_'+scope+'_'+str(selected)+'_'+stage+'_'+label)
                inputs[stage]=values_out
        st.caption('같은 기성의 분할 수령은 합산합니다. 기성별 보험료·제외항목을 입력하고 관리비를 원 미만 버림한 뒤 합산합니다.')
        preview=st.form_submit_button('관리비 계산 확인')
        save=st.form_submit_button('관리비·보험료 정산 저장',type='primary')
    if preview:
        lines=flow.fee_lines(c,p,m['id'],ids,inputs)
        display=[]
        for v in lines:
            excluded=sum((db.num(v[k]) for k in db.EXCLUSIONS),db.num(0))
            display.append({'기성':v['label'],'선택 공급가액':money(v['supply']),'제외분 합계':money(excluded),'관리대상액':money(db.num(v['supply'])-excluded),'관리비율':v['fee_rate']+'%','관리비':money(v['fee'])})
        st.dataframe(display,hide_index=True,use_container_width=True)
        st.success('관리비 합계: '+money(sum((db.num(v['fee']) for v in lines),db.num(0)))+'원')
    if save:
        gid=flow.save_group(c,p,m['id'],name,ids,inputs,selected);st.session_state['_jv_next_fee']=[p,m['id'],gid];saved()
    if selected and st.checkbox('선택한 관리비 정산을 삭제합니다.',key='fee_delete_'+scope):
        if st.button('관리비 정산 삭제'):flow.delete_group(c,p,m['id'],selected);saved()

def entry_page(c,p,m):
    st.header('공사대금·관리비·하도급·회수')
    tabs=st.tabs(['공사대금 입력','관리비 및 보험료 입력','하도급 청구 및 지급현황 입력','회수내역 입력'])
    with tabs[0]:entry_form(c,p,m,'receipt')
    with tabs[1]:fee_page(c,p,m)
    with tabs[2]:entry_form(c,p,m,'subcontract')
    with tabs[3]:entry_form(c,p,m,'recovery')

def issue_settings(p,m,kind,data,month=None):
    owner=next(x for x in data['companies'] if x['id']==data['project']['owner_id']);info=json.loads(owner['info']);scope=f"{p}_{m['id']}_{kind}_{month or 'run'}"
    with st.expander('발행 정보 · 문서번호 수동입력 · 공문 문구',expanded=False):
        cols=st.columns(2);number=cols[0].text_input('문서번호 (수동입력)',key='number_'+scope);issued=cols[1].date_input('발행일',key='issued_'+scope)
        meta=dict(number=number,issued=issued.isoformat())
        if kind=='공문':
            meta.update(title=st.text_input('제목',period_label(month,month)+' 공동원가 안분정산 및 자금청구의 건',key='title_'+scope),reference=st.text_input('참조','공동도급 담당자',key='reference_'+scope),body=st.text_area('본문 추가 문구',key='body_'+scope),due=st.text_input('입금 요청기한',key='due_'+scope),bank=st.text_input('입금계좌',info.get('입금계좌',''),key='bank_'+scope),attachments=st.text_input('첨부목록','공동원가 투입 및 배분내역서 1부, 현장별 원장 1부',key='attachments_'+scope))
            seal=st.file_uploader('직인 이미지 (선택, PNG/JPG 1MB 이하)',type=['png','jpg','jpeg'],key='seal_'+scope)
            if seal:
                import base64
                content=seal.getvalue();mime='image/png' if content.startswith(b'\x89PNG\r\n\x1a\n') else 'image/jpeg' if content.startswith(b'\xff\xd8\xff') else None
                if len(content)>1024*1024 or not mime:raise ValueError('1MB 이하 PNG/JPG 직인을 선택하세요.')
                meta['seal']='data:'+mime+';base64,'+base64.b64encode(content).decode('ascii')
    return number,issued,meta

def settlement_page(c,p,m):
    st.header('원가정산서 · 내역 선택')
    rs=flow.runs(c,p,m['id'])
    if st.session_state.get('_jv_next_run',[])[:2]==[p,m['id']]:st.session_state[f"run_{p}_{m['id']}"]=st.session_state.pop('_jv_next_run')[2]
    if st.session_state.get(f"run_{p}_{m['id']}") not in [None]+[r['id'] for r in rs]:st.session_state[f"run_{p}_{m['id']}"]=None
    selected=st.selectbox('정산 건',[None]+[r['id'] for r in rs],format_func=lambda id:'새 정산' if id is None else next(r['name'] for r in rs if r['id']==id),key=f"run_{p}_{m['id']}")
    old=next((r for r in rs if r['id']==selected),{});default=json.loads(old.get('selection','{}'));scope=f"{p}_{m['id']}_{selected}"
    name=st.text_input('정산명',old.get('name',''),key='run_name_'+scope);selection={}
    source=flow.source_entries(c,p,m['id'],'receipt')
    selection['receipts']=pick_table('① 공사대금 수령 범위',[{'id':e['id'],'일자':e['entry_date'],'기성':e['data']['label'],'공급가액':e['data']['supply'],'VAT':e['data']['vat']} for e in source],default.get('receipts',[]),'run_receipts_'+scope,flow.used(c,p,m['id'],'receipts',selected))
    selection['fees']=pick_table('② 관리비·보험료 정산 범위',[{'id':g['id'],'정산명':g['label'],'수령 건수':len(json.loads(g['receipt_ids']))} for g in flow.groups(c,p,m['id'])],default.get('fees',[]),'run_fees_'+scope,flow.used(c,p,m['id'],'fees',selected))
    claims=flow.claim_rows(c,p,m['id'])
    months=st.multiselect('청구 범위 일괄 선택 (배분월)',sorted({r['month'] for r in claims}),key='bulk_months_'+scope)
    if st.button('선택한 월의 청구 내역 적용',key='bulk_apply_'+scope):
        st.session_state['claims_override_'+scope]=[r['id'] for r in claims if r['month'] in months]
        st.session_state['claims_bulk_revision_'+scope]=st.session_state.get('claims_bulk_revision_'+scope,0)+1
    claim_defaults=st.session_state.get('claims_override_'+scope,default.get('claims',[]))
    selection['claims']=pick_table('③ 원가분담금 청구 범위',[{'id':r['id'],'배분월':r['month'],'계정과목':r['account'],'적요':r['description'],'공급가액':r['base'],'VAT':r['vat'],'회원사투입 차감':r['invested']} for r in claims],claim_defaults,'run_claims_'+scope+'_'+str(st.session_state.get('claims_bulk_revision_'+scope,0)),flow.used(c,p,m['id'],'claims',selected))
    for kind,key,title in [('subcontract','subcontracts','④ 하도급 청구·지급 범위'),('recovery','recoveries','⑤ 회수 범위')]:
        selection[key]=pick_table(title,[{'id':e['id'],'일자':e['entry_date'],'업체·비고':e['data'].get('vendor',e['data'].get('note','')),'공급가액':e['data']['supply'],'VAT':e['data']['vat']} for e in flow.source_entries(c,p,m['id'],kind)],default.get(key,[]),'run_'+key+'_'+scope,flow.used(c,p,m['id'],key,selected))
    if st.button('선택 내역으로 정산 저장',type='primary'):
        rid=flow.save_run(c,p,m['id'],name,selection,selected);st.session_state.pop('claims_override_'+scope,None);st.session_state['_jv_next_run']=[p,m['id'],rid];saved()
    if selected:
        with st.expander('정산 연결 해제'):
            confirmed=st.checkbox('이 정산의 선택 연결을 삭제합니다. 공사대금·관리비·원장·발행본은 유지됩니다.',key='run_delete_'+scope)
            if st.button('정산 건 연결 삭제',disabled=not confirmed):flow.delete_run(c,p,m['id'],selected);saved()
    if not any(selection.values()):st.info('기간이 아닌, 이번 정산에 포함할 내역을 각각 선택하세요.');return
    try:data=flow.selected_settlement(c,p,m['id'],selection)
    except ValueError as e:st.warning(str(e));return
    data['run_name']=name
    number,issued,meta=issue_settings(p,m,'정산서',data)
    html=report_html(data,'정산서',meta)
    st.download_button('선택 정산서 HTML 내려받기',html.encode('utf-8'),file_name='원가정산서.html',mime='text/html')
    if st.button('현재 정산서 확정·발행'):
        if not selected or json.loads(old['selection'])!=data['selection']:raise ValueError('선택 내역을 먼저 정산 저장한 뒤 발행하세요.')
        db.snapshot(c,data,'정산서',number,issued,meta);saved()
    components.html(html,height=1900,scrolling=True)

def monthly_reports(c,p,m,user,letter_only=False):
    st.header('공문 작성' if letter_only else '월 마감 · 보고서 출력')
    month=st.text_input('배분연월 (YYYY-MM)',date.today().strftime('%Y-%m'),key='letter_month' if letter_only else 'close_month');db.period(month)
    closed=flow.is_closed(c,p,month)
    st.success('마감 완료 · 보고서 출력 가능') if closed else st.info('미마감 · 원장과 원가배분 확인 후 마감하면 출력할 수 있습니다.')
    if not letter_only:
        if st.button('선택 월 마감',type='primary',disabled=closed):flow.close_month(c,p,month,st.session_state['_auth_token']);saved()
        if user['admin'] and closed:
            reason=st.text_input('마감 해제 사유')
            if st.button('관리자 마감 해제'):flow.reopen(c,p,month,st.session_state['_auth_token'],reason);saved()
    if letter_only and not closed:
        with st.expander('공문 전월 미수액 · 지연이자 입력'):
            entry_form(c,p,m,'opening',month)
    data=flow.closed_package(c,p,month)['letters'][str(m['id'])] if closed else flow.month_letter(c,p,m['id'],month)
    number,issued,meta=issue_settings(p,m,'공문',data,month)
    html=report_html(data,'공문',meta,printable=closed)
    st.download_button('공문 HTML 내려받기',html.encode('utf-8'),file_name='공문_'+month+'.html',mime='text/html',disabled=not closed)
    if st.button('공문 확정·발행',disabled=not closed):
        fresh=flow.closed_package(c,p,month)['letters'][str(m['id'])];db.snapshot(c,fresh,'공문',number,issued,meta);saved()
    if closed:
        from joint_report import monthly_package
        package=flow.closed_package(c,p,month);bundle,ledger_html,alloc_html=monthly_package(package,m['id'],meta)
        st.download_button('보고서 출력 묶음 · 공문/해당월 원장/원가배분내역서',bundle,file_name='공동도급_'+month+'_출력.zip',mime='application/zip')
        a,b=st.columns(2)
        a.download_button('해당월 원장 출력 HTML',ledger_html.encode('utf-8'),file_name='원장_'+month+'.html',mime='text/html')
        b.download_button('원가배분내역서 출력 HTML',alloc_html.encode('utf-8'),file_name='원가배분_'+month+'.html',mime='text/html')
        st.caption('출력 묶음에는 보관된 업로드 원본도 포함됩니다. 이전 버전에서 원본을 보관하지 않은 건은 해당월 원장 출력표로 제공합니다.')
    else:st.button('보고서 출력 묶음',disabled=True)
    if letter_only:components.html(html,height=1600,scrolling=True)

def archive_page(c,p):
    st.header('발행본·백업·변경 이력')
    docs=db.rows(c,'SELECT id,kind,number,issued FROM jv_documents WHERE project_id=? ORDER BY id DESC',(p,));st.dataframe(docs,hide_index=True)
    if docs:
        id=st.selectbox('확정 발행본',[x['id'] for x in docs],format_func=lambda id:next(x['number']+' · '+x['kind'] for x in docs if x['id']==id))
        d=json.loads(c.execute('SELECT snapshot FROM jv_documents WHERE id=?',(id,)).fetchone()[0]);d['snapshot_id']=id
        html=report_html(d,d['kind']);st.download_button('확정본 내려받기',html.encode('utf-8'),file_name=d['kind']+'_'+str(id)+'.html',mime='text/html');components.html(html,height=800,scrolling=True)
    st.download_button('변경 이력 CSV',safe_csv(pd.DataFrame(db.rows(c,'SELECT at,actor,table_name,operation,before_json,after_json FROM change_history ORDER BY id DESC LIMIT 1000'))),file_name='공동도급_변경이력.csv')
    if st.button('현재 DB 백업 생성'):
        path=shared.backup_database(c.path);st.download_button('생성한 백업 내려받기',path.read_bytes(),file_name=path.name)
    st.caption('매일 첫 변경 전 자동 백업합니다. 복원은 관리자만 실행할 수 있습니다.')
    if shared.validate_session(st.session_state['_auth_token'])['admin']:
        up=st.file_uploader('공동도급 DB 백업 복원',type=['sqlite','db'],key='jv_restore')
        confirm=st.checkbox('공동도급 자료 전체를 선택한 백업으로 되돌립니다.')
        if st.button('공동도급 백업 복원 실행',disabled=not(up and confirm)):
            db.restore(c,up.getvalue(),st.session_state['_auth_token']);st.session_state.pop('_snapshots',None);saved()

def render(user):
    st.caption('공동도급 관리대장 · 버전 '+VERSION)
    c=db.connect()
    with st.sidebar:
        st.subheader('공동도급 관리대장')
        page=st.radio('공동도급 업무',['법인·담당자','프로젝트·회원사','원장 등록','원가배분내역서','기성·하도급·회수','원가정산서','공문 작성','월 마감·보고서 출력','발행본·백업'],key='jv_page')
    try:
        if '_jv_success' in st.session_state:st.success(st.session_state.pop('_jv_success'))
        ps=db.projects(c);p=None;m=None
        if ps:
            cols=st.columns([2,2,1])
            p=cols[0].selectbox('공동도급 프로젝트',[x['id'] for x in ps],format_func=lambda id:next(x['name'] for x in ps if x['id']==id),key='jv_project')
            ms=db.members(c,p);options=[x['id'] for x in ms]
            if st.session_state.get('jv_member') not in options:st.session_state['jv_member']=options[0] if options else None
            selected=cols[1].selectbox('대상 회원사 (프로젝트 등록 회사)',options,format_func=lambda id:next(x['name'] for x in ms if x['id']==id),key='jv_member')
            m=next((x for x in ms if x['id']==selected),None)
            if m:cols[2].caption('지분 '+m['share']+'% · 관리비 '+(m['fee_rate']+'%' if m['fee_enabled'] else '미적용'))
        if page=='법인·담당자':company_page(c);return
        if page=='프로젝트·회원사':project_page(c,p);return
        if p is None or not m:st.info('법인 등록 후 프로젝트·회원사를 등록하세요.');return
        if page=='원장 등록':ledger_page(c,p,ms)
        elif page=='발행본·백업':archive_page(c,p)
        elif page=='기성·하도급·회수':entry_page(c,p,m)
        elif page=='원가정산서':settlement_page(c,p,m)
        elif page=='공문 작성':monthly_reports(c,p,m,user,True)
        elif page=='월 마감·보고서 출력':monthly_reports(c,p,m,user)
        elif page=='원가배분내역서':
            month=st.text_input('배분연월 (YYYY-MM)',date.today().strftime('%Y-%m'));db.period(month);allocation_page(c,p,month,month)
    except (ValueError,sqlite3.Error,BadZipFile) as e:st.error(str(e))
