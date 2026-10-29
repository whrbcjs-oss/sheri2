"""Read the classified ordinary-bank ledger; no report worksheet is required."""
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import re
import unicodedata
from pathlib import Path
import openpyxl
from monthly_report_preview import WorkbookCalculator
from openpyxl.formula.tokenizer import Tokenizer

LABELS={8:'공사수입',9:'김제공장',10:'입금 계 ①',11:'급여 등',12:'소득세등',13:'노임',
        15:'건강.연금',16:'고용.산재',17:'미불',19:'수시',20:'공과금',21:'카드대금',
        22:'이자비용',23:'지출 계 ②',24:'수입·지출 차액 ①−②',25:'관리비',26:'관리비차감후'}
EXPENSE=[11,12,13,15,16,17,19,20,21,22]
HEADERS={'date':['날짜','거래일자'],'inflow':['차변','입금'],'outflow':['대변','출금'],
         'department':['대분류'],'category':['중분류'],'vat':['VAT','부가세'],
         'corporation':['법인'],'project':['프로젝트명','프로젝트명2'],'vendor':['거래처명'],
         'description':['품의내역명','적요란','적요'],'note':['비고']}


class LedgerValueReader(WorkbookCalculator):
    """Recalculate supported formulas; use saved Excel results for other formulas.

    This reader is only for importing the ledger. The separate independent report
    verifier continues to reject unsupported formulas instead of trusting caches.
    """
    FUNCTIONS={'SUMIFS','SUM','MONTH'}
    TEXT_FUNCTIONS={'IF','IFERROR','OR','AND','NOT','ISNUMBER','SEARCH'}

    def __init__(self,book,cached):
        super().__init__(book)
        self.cached=cached
        self.cached_cells=set()
        self.native_cells=set()
        self._values={}
        self._active=set()

    def value(self,sheet,coordinate):
        key=(sheet,coordinate)
        if key in self._values:return self._values[key]
        if key in self._active:raise ValueError(f'{sheet}!{coordinate}: 순환 참조 수식입니다.')
        self._active.add(key)
        try:
            result=self._read_value(sheet,coordinate)
            self._values[key]=result
            return result
        finally:self._active.remove(key)

    def _read_value(self,sheet,coordinate):
        cell=self.book[sheet][coordinate]
        if cell.data_type=='e':raise ValueError(f'{sheet}!{coordinate}: 엑셀 오류 {cell.value}가 있습니다.')
        if cell.data_type!='f':return cell.value
        supported=isinstance(cell.value,str)
        functions=[]
        tokens=[]
        if supported:
            tokens=[t for t in Tokenizer(cell.value).items if t.type!='WHITE-SPACE']
            functions=[t.value[:-1].upper() for t in tokens if t.type=='FUNC' and t.subtype=='OPEN']
            supported=all(name in self.FUNCTIONS for name in functions)
            supported=supported and all(t.value in ('+','-','*','/') for t in tokens if t.type=='OPERATOR-INFIX')
            supported=supported and all(t.subtype in ('NUMBER','TEXT','RANGE') for t in tokens if t.type=='OPERAND')
        if supported:
            return self.expression(cell.value,sheet)
        if tokens and self.supports_text_formula(tokens):
            result=self.text_formula(tokens,sheet)
            if isinstance(result,FormulaError):
                raise ValueError(f'{sheet}!{coordinate}: 분류 수식 계산 오류 {result.code}입니다. 원본 수식을 확인하세요.')
            self.native_cells.add(f'{sheet}!{coordinate}')
            return result
        cached=self.cached[sheet][coordinate]
        if cached.data_type=='e':
            raise ValueError(f'{sheet}!{coordinate}: 엑셀 계산 결과에 오류 {cached.value}가 있습니다. 원본 수식을 확인하고 다시 계산·저장하세요.')
        # A string-typed empty cached result is distinct from a missing numeric cache.
        if cached.value is not None or cached.data_type=='str':
            self.cached_cells.add(f'{sheet}!{coordinate}')
            return '' if cached.value is None else cached.value
        names=', '.join(dict.fromkeys(name for name in functions if name not in self.FUNCTIONS)) or '해당'
        raise ValueError(f'{sheet}!{coordinate}: {names} 수식의 저장된 계산 결과가 없습니다. '
                         'Excel에서 계산 옵션을 자동으로 설정하고 Ctrl+Alt+F9로 전체 계산한 뒤 Ctrl+S로 저장하여 다시 업로드하세요.')


    def supports_text_formula(self,tokens):
        for token in tokens:
            if token.type=='FUNC' and token.subtype=='OPEN' and token.value[:-1].upper() not in self.TEXT_FUNCTIONS:return False
            if token.type not in ('FUNC','SEP','OPERAND'):return False
            if token.type=='OPERAND' and token.subtype not in ('NUMBER','TEXT','RANGE','LOGICAL'):return False
            if token.type=='OPERAND' and token.subtype=='RANGE' and (':' in token.value or '[[' in token.value):return False
        return True

    def text_formula(self,tokens,sheet):
        pos=0
        def parse():
            nonlocal pos
            token=tokens[pos];pos+=1
            if token.type=='FUNC' and token.subtype=='OPEN':
                args=[]
                while pos<len(tokens) and tokens[pos].subtype!='CLOSE':
                    args.append(parse())
                    if pos<len(tokens) and tokens[pos].type=='SEP':pos+=1
                    else:break
                if pos>=len(tokens) or tokens[pos].subtype!='CLOSE':raise ValueError('분류 수식 괄호를 확인하세요.')
                pos+=1;return ('call',token.value[:-1].upper(),args)
            if token.type!='OPERAND':raise ValueError('지원하지 않는 분류 수식입니다.')
            return ('operand',token.subtype,token.value)
        tree=parse()
        if pos!=len(tokens):raise ValueError('지원하지 않는 분류 수식입니다.')
        def truth(value):
            if isinstance(value,FormulaError):return value
            if value is None:return False
            if isinstance(value,(bool,int,float,Decimal)):return bool(value)
            if str(value).upper() in ('TRUE','FALSE'):return str(value).upper()=='TRUE'
            return FormulaError('#VALUE!')
        def evaluate(node):
            kind,name,args=node
            if kind=='operand':
                if name=='NUMBER':return Decimal(args)
                if name=='TEXT':return args[1:-1].replace('""','"')
                if name=='LOGICAL':return args.upper()=='TRUE'
                return self.resolve(args,sheet)
            if name=='IF':
                if len(args) not in (2,3):return FormulaError('#VALUE!')
                condition=truth(evaluate(args[0]))
                if isinstance(condition,FormulaError):return condition
                return evaluate(args[1]) if condition else evaluate(args[2]) if len(args)==3 else False
            if name=='IFERROR':
                if len(args)!=2:return FormulaError('#VALUE!')
                result=evaluate(args[0]);return evaluate(args[1]) if isinstance(result,FormulaError) else result
            values=[evaluate(arg) for arg in args]
            if name=='ISNUMBER':
                if len(values)!=1:return FormulaError('#VALUE!')
                return isinstance(values[0],(Decimal,int,float)) and not isinstance(values[0],bool)
            for value in values:
                if isinstance(value,FormulaError):return value
            if name in ('OR','AND','NOT'):
                converted=[truth(value) for value in values]
                for value in converted:
                    if isinstance(value,FormulaError):return value
                if name=='OR':return any(converted)
                if name=='AND':return all(converted)
                return not converted[0] if len(converted)==1 else FormulaError('#VALUE!')
            if name=='SEARCH':
                if len(values) not in (2,3):return FormulaError('#VALUE!')
                find='' if values[0] is None else str(values[0])
                text='' if values[1] is None else str(values[1])
                try:start=int(values[2]) if len(values)==3 else 1
                except (TypeError,ValueError):return FormulaError('#VALUE!')
                if start<1 or start>len(text)+1:return FormulaError('#VALUE!')
                pattern=[];i=0
                while i<len(find):
                    char=find[i]
                    if char=='~' and i+1<len(find):i+=1;pattern.append(re.escape(find[i]))
                    elif char=='*':pattern.append('.*')
                    elif char=='?':pattern.append('.')
                    else:pattern.append(re.escape(char))
                    i+=1
                match=re.search(''.join(pattern),text[start-1:],flags=re.IGNORECASE|re.DOTALL)
                return Decimal(match.start()+start) if match else FormulaError('#VALUE!')
            return FormulaError('#NAME?')
        return evaluate(tree)


class FormulaError:
    def __init__(self,code):self.code=code


def money(value):
    if value is None or value=='':return Decimal(0)
    text=str(value)
    start,end=0,len(text)
    def padding(char):return char.isspace() or unicodedata.category(char)=='Cf'
    while start<end and padding(text[start]):start+=1
    while end>start and padding(text[end-1]):end-=1
    text=text[start:end]
    if not text:return Decimal(0)
    text=text.replace(',','')
    if text.startswith('(') and text.endswith(')'):text='-'+text[1:-1]
    try:amount=Decimal(text)
    except InvalidOperation:raise ValueError(f'금액을 해석할 수 없습니다: {value!r}')
    if not amount.is_finite():raise ValueError('유한한 금액을 입력하세요.')
    return amount


def parse_date(value,epoch):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    if isinstance(value,(int,float)) and 30000<value<80000:
        return openpyxl.utils.datetime.from_excel(value,epoch).date()
    text=str(value).strip()[:10]
    for fmt in ['%Y-%m-%d','%Y.%m.%d','%Y/%m/%d','%Y%m%d']:
        try:return datetime.strptime(text,fmt).date()
        except ValueError:pass
    raise ValueError('날짜 형식을 확인하세요: '+str(value))


def read_ledger(path):
    path=Path(path);book=openpyxl.load_workbook(path,data_only=False)
    cached=openpyxl.load_workbook(path,data_only=True)
    sheets=[book['보통예금_원장']] if '보통예금_원장' in book.sheetnames else list(book)
    selected=None
    for sheet in sheets:
        for row in sheet.iter_rows(min_row=1,max_row=min(25,sheet.max_row)):
            names={str(c.value).strip():c.column for c in row if c.value is not None}
            mapping={key:next((names[n] for n in aliases if n in names),None) for key,aliases in HEADERS.items()}
            if all(mapping[k] for k in ['date','inflow','outflow','department','category','vat']):
                if selected is not None:raise ValueError('원장 형식의 시트가 여러 개입니다. 보통예금_원장 이름으로 저장해 주세요.')
                selected=(sheet,row[0].row,mapping);break
    if selected is None:raise ValueError('날짜, 차변/입금, 대변/출금, 대분류, 중분류, VAT 열이 있는 보통예금 원장이 필요합니다.')
    sheet,header,mapping=selected;calc=LedgerValueReader(book,cached);records=[]
    for r in range(header+1,sheet.max_row+1):
        def value(key):
            column=mapping.get(key)
            return calc.value(sheet.title,f'{openpyxl.utils.get_column_letter(column)}{r}') if column else None
        def amount(key):
            column=mapping[key]
            coordinate=f'{openpyxl.utils.get_column_letter(column)}{r}'
            try:return money(value(key))
            except (ValueError,TypeError,ArithmeticError) as e:
                label=str(sheet.cell(header,column).value)
                raise ValueError(f'{sheet.title}!{coordinate} ({label}): {e}') from e
        raw_date=value('date')
        if raw_date in (None,''):
            # Blank rows and unclassified footer totals are not transactions.
            if str(value('department') or '')=='데크거더' and any(amount(k) for k in ['inflow','outflow','vat']):
                raise ValueError(f'{r}행: 집계 대상 금액에 거래일자가 없습니다.')
            continue
        try:
            posting=parse_date(raw_date,book.epoch)
            if not 2000<=posting.year<=2100:raise ValueError('지원 범위를 벗어난 거래 연도입니다.')
            entry={'row':r,'date':posting.isoformat(),'year':posting.year,'month':posting.month}
            for key in ['inflow','outflow','vat']:entry[key]=str(amount(key))
            for key in ['department','category','corporation','project','vendor','description','note']:
                entry[key]=str(value(key) or '').strip()
            records.append(entry)
        except (ValueError,TypeError,ArithmeticError) as e:raise ValueError(f'{r}행: {e}') from e
    if not records:raise ValueError('읽을 수 있는 거래가 없습니다.')
    return {'source':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'sheet':sheet.title,'records':records,'cached_formula_count':len(calc.cached_cells),'native_formula_count':len(calc.native_cells),
            'calculation_mode':getattr(book.calculation,'calcMode',None),
            'periods':sorted({(r['year'],r['month']) for r in records})}


def build_report(year,records,imports,settings,adjustments):
    """Calculate from stored monthly records, with explicit non-ledger inputs."""
    months=sorted(i['month'] for i in imports)
    if not months:return None
    selected=[r for r in records if r['department']=='데크거더']
    matrix={r:[Decimal(0)]*12 for r in LABELS};breakdowns=[];applied=[]
    for month in months:
        for row in [8,9]+EXPENSE:
            category='공사수입' if row==8 else '제품매출' if row==9 else LABELS[row]
            group=[r for r in selected if r['month']==month and r['category']==category]
            gross=sum((money(r['inflow'] if row in (8,9) else r['outflow']) for r in group),Decimal(0))
            vat=sum((money(r['vat']) for r in group),Decimal(0))
            matrix[row][month-1]=gross-vat
            breakdowns.append({'report_row':row,'month':month,'gross':float(gross),'vat':float(vat),'count':len(group)})
        for item in adjustments:
            if item['month']==month:
                amount=money(item['amount']);row=item['report_row'];matrix[row][month-1]+=amount
                applied.append({'row':row,'month':month,'amount':float(amount),'cell':f'{month}월 {LABELS[row]}',
                                'terms':[str(amount)],'formula':'별도 입력: '+item['reason']})
        m=month-1
        matrix[10][m]=matrix[8][m]+matrix[9][m]
        matrix[23][m]=sum((matrix[r][m] for r in EXPENSE),Decimal(0))
        matrix[24][m]=matrix[10][m]-matrix[23][m]
        matrix[25][m]=matrix[10][m]*money(settings['management_rate'])/100
        matrix[26][m]=matrix[24][m]-matrix[25][m]
    opening=money(settings['opening_amount'])
    rate=str(settings['management_rate'])
    notes={11:'퇴직연금·공제 등 실제 지급액',12:'근로·사업·지방소득세',15:'건강·연금 등 원장 분류 기준',
           16:'원청 공제 보험료는 원장외 조정으로 별도 입력',20:'원장 공과금 분류 기준',21:'원장 카드대금 분류 기준',
           22:'원장 이자비용 분류 기준',25:'입금 합계 × '+rate+'%'}
    rows=[{'row':r,'label':f'관리비({rate}%)' if r==25 else LABELS[r],
           'values':[float(v) for v in matrix[r]]+[float(sum(matrix[r],Decimal(0)))],
           'note':notes.get(r,''),'subnotes':[]} for r in LABELS]
    refunds=[{**r,**{k:float(money(r[k])) for k in ['inflow','outflow','vat']}} for r in selected if '환수' in r['category']]
    return {'validation_mode':'ledger','source':f'{year}년 월별 업로드 원장 ({len(months)}개월)',
            'sheet':'보통예금_원장 → 자동 생성','period_start':min(r['date'] for r in records),
            'period_end':max(r['date'] for r in records),'months':months,'year':year,'management_rate':float(money(rate)),
            'ledger_count':len(records),'selected_count':len(selected),'month_counts':dict(Counter(r['month'] for r in selected)),
            'rows':rows,'breakdowns':breakdowns,'checks':[], 'adjustments':applied,'refunds':refunds,
            'refund_total':sum(r['inflow'] for r in refunds),'vat_scope_impact':0,
            'opening':float(opening),'opening_confirmed':bool(settings['opening_confirmed']),
            'closing':float(opening+sum(matrix[26],Decimal(0))) if settings['opening_confirmed'] else None,
            'prior_year_claims':None,'imports':imports,'notes':[]}
