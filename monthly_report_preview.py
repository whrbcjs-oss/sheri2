"""Read-only workbook analysis and a standalone monthly-report design preview.

The original workbook and all generated financial data remain on the user's PC.
Only this source file belongs in the public repository.
"""
from __future__ import annotations
import argparse
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from functools import lru_cache
import hashlib
import json
from pathlib import Path
import re
import webbrowser

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

REPORT = '수입.지출(총)(부가세제외)'
LEDGER = '보통예금_원장'
NUMERIC_ROWS = [8,9,10,11,12,13,15,16,17,19,20,21,22,23,24,25,26]
EXPENSE_ROWS = [11,12,13,15,16,17,19,20,21,22]


def number(value):
    if value is None or value == '':
        return Decimal(0)
    if isinstance(value, (Decimal,int,float)) and not isinstance(value,bool):
        return Decimal(str(value))
    raise ValueError('금액 자리에 숫자가 아닌 값이 있습니다: ' + str(value))


@dataclass
class Range:
    sheet: str
    coordinates: list[str]


class WorkbookCalculator:
    """Restricted Excel evaluator: SUMIFS, SUM, MONTH, arithmetic, references.

    No eval/exec and no workbook macros, links, or document instructions are run.
    Unsupported formulas fail explicitly rather than silently using cached data.
    """
    def __init__(self, book):
        self.book = book

    def resolve(self, text, sheet):
        if '!' in text:
            sheet, text = text.rsplit('!',1)
            sheet = sheet.strip("'").replace("''", "'")
        if '[[' in text:
            table_name = text.split('[[',1)[0]
            header = re.findall(r'\[([^\[\]]+)\]',text)[-1]
            for ws in self.book:
                for table in ws.tables.values():
                    if table.name == table_name:
                        start,end = table.ref.split(':')
                        c1,r1 = openpyxl.utils.cell.coordinate_from_string(start)
                        c2,r2 = openpyxl.utils.cell.coordinate_from_string(end)
                        headers = {ws.cell(r1,c).value:c for c in range(column_index_from_string(c1),column_index_from_string(c2)+1)}
                        col = headers[header]
                        return Range(ws.title,[f'{get_column_letter(col)}{r}' for r in range(r1,r2+1)])
            raise ValueError('표를 찾을 수 없습니다: '+table_name)
        text=text.replace('$','')
        if ':' not in text:
            return self.value(sheet,text)
        start,end=text.split(':')
        if start.isalpha() and end.isalpha():
            if start!=end:raise ValueError('여러 열 전체 참조는 지원하지 않습니다.')
            return Range(sheet,[f'{start}{r}' for r in range(1,self.book[sheet].max_row+1)])
        a,b=openpyxl.utils.cell.coordinate_from_string(start),openpyxl.utils.cell.coordinate_from_string(end)
        return Range(sheet,[f'{get_column_letter(c)}{r}' for r in range(a[1],b[1]+1) for c in range(column_index_from_string(a[0]),column_index_from_string(b[0])+1)])

    @lru_cache(maxsize=None)
    def value(self,sheet,coordinate):
        c=self.book[sheet][coordinate]
        if c.data_type=='f':return self.expression(c.value,sheet)
        if c.data_type=='e':raise ValueError(f'원본 셀 오류: {sheet}!{coordinate} {c.value}')
        return c.value

    def expression(self,formula,sheet):
        tokens=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE']
        pos=0
        def primary():
            nonlocal pos
            token=tokens[pos];pos+=1
            if token.type=='OPERATOR-PREFIX':
                result=number(primary());return -result if token.value=='-' else result
            if token.type=='PAREN' and token.subtype=='OPEN':
                result=expr();assert tokens[pos].subtype=='CLOSE';pos+=1
            elif token.type=='FUNC' and token.subtype=='OPEN':
                args=[]
                while tokens[pos].subtype!='CLOSE':
                    args.append(expr())
                    if tokens[pos].type=='SEP':pos+=1
                    else:break
                if tokens[pos].subtype!='CLOSE':raise ValueError('함수 닫기 오류')
                pos+=1
                result=self.function(token.value[:-1],args)
            elif token.type=='OPERAND':
                if token.subtype=='NUMBER':result=Decimal(token.value)
                elif token.subtype=='TEXT':result=token.value[1:-1].replace('""','"')
                elif token.subtype=='RANGE':result=self.resolve(token.value,sheet)
                else:raise ValueError('지원하지 않는 수식 토큰: '+token.value)
            else:raise ValueError('지원하지 않는 수식 토큰: '+token.value)
            while pos<len(tokens) and tokens[pos].type=='OPERATOR-POSTFIX':
                if tokens[pos].value!='%':raise ValueError('지원하지 않는 후위 연산자')
                result=number(result)/100;pos+=1
            return result
        def expr(minimum=0):
            nonlocal pos
            left=primary();precedence={'+':1,'-':1,'*':2,'/':2}
            while pos<len(tokens) and tokens[pos].type=='OPERATOR-INFIX':
                op=tokens[pos].value
                if op not in precedence:raise ValueError('지원하지 않는 연산자: '+op)
                power=precedence[op]
                if power<minimum:break
                pos+=1;right=expr(power+1)
                a,b=number(left),number(right)
                if op=='+':left=a+b
                elif op=='-':left=a-b
                elif op=='*':left=a*b
                else:left=a/b
            return left
        result=expr()
        if pos!=len(tokens):raise ValueError('해석하지 못한 수식: '+formula)
        return result

    def function(self,name,args):
        name=name.upper()
        if name=='MONTH':
            value=args[0]
            if hasattr(value,'month'):return Decimal(value.month)
            match=re.match(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})',str(value))
            if not match:raise ValueError('날짜를 해석할 수 없습니다: '+str(value))
            return Decimal(match.group(2))
        if name=='SUM':
            total=Decimal(0)
            for arg in args:
                values=[self.value(arg.sheet,c) for c in arg.coordinates] if isinstance(arg,Range) else [arg]
                total+=sum((number(v) for v in values if isinstance(v,(int,float,Decimal)) and not isinstance(v,bool)),Decimal(0))
            return total
        if name=='SUMIFS':
            source=args[0];criteria=list(zip(args[1::2],args[2::2]));total=Decimal(0)
            if not isinstance(source,Range) or any(not isinstance(r,Range) or len(r.coordinates)!=len(source.coordinates) for r,_ in criteria):
                raise ValueError('SUMIFS 범위 길이가 일치하지 않습니다.')
            for i,coord in enumerate(source.coordinates):
                if all(self.matches(self.value(r.sheet,r.coordinates[i]),wanted) for r,wanted in criteria):
                    value=self.value(source.sheet,coord)
                    if isinstance(value,(Decimal,int,float)):total+=number(value)
            return total
        raise ValueError('지원하지 않는 함수: '+name)

    @staticmethod
    def matches(value,wanted):
        if isinstance(wanted,(int,float,Decimal)):
            return isinstance(value,(int,float,Decimal)) and number(value)==number(wanted)
        if isinstance(wanted,str) and (wanted.startswith(('>','<','=')) or any(c in wanted for c in '*?~')):
            raise ValueError('지원하지 않는 SUMIFS 비교 조건: '+wanted)
        return str(value or '').casefold()==str(wanted or '').casefold()


def analyze(path, include_ledger=False):
    path=Path(path)
    formulas=openpyxl.load_workbook(path,data_only=False)
    saved=openpyxl.load_workbook(path,data_only=True)
    for name in [REPORT,LEDGER]:
        if name not in formulas.sheetnames:raise ValueError('필수 시트가 없습니다: '+name)
    calc=WorkbookCalculator(formulas);ws=saved[REPORT]
    rows=[];checks=[]
    for r in NUMERIC_ROWS:
        values=[]
        for col in range(3,16):
            coordinate=f'{get_column_letter(col)}{r}'
            expected=number(calc.value(REPORT,coordinate))
            cached=ws[coordinate].value
            delta=expected-number(cached) if isinstance(cached,(int,float)) else None
            checks.append({'cell':coordinate,'cached':cached,'recalculated':float(expected),'difference':float(delta) if delta is not None else None,'ok':delta is not None and abs(delta)<Decimal('0.01')})
            values.append(float(expected))
        rows.append({'row':r,'label':str(ws.cell(r,2).value).strip(),'values':values,'note':str(ws.cell(r,16).value or ''),'subnotes':[str(ws.cell(r+1,c).value or '') for c in range(3,15)] if r in (13,17) else []})
    ledger=[]
    for r in range(4,formulas[LEDGER].max_row+1):
        date=calc.value(LEDGER,f'C{r}')
        if date is None:continue
        ledger.append({'row':r,'date':str(date)[:10],'month':int(number(calc.value(LEDGER,f'D{r}'))),
                       'department':calc.value(LEDGER,f'Q{r}'),'category':calc.value(LEDGER,f'R{r}'),
                       'inflow':float(number(calc.value(LEDGER,f'K{r}'))),'outflow':float(number(calc.value(LEDGER,f'L{r}'))),
                       'vat':float(number(calc.value(LEDGER,f'T{r}'))),
                       **({key:str(calc.value(LEDGER,f'{col}{r}') or '') for key,col in [('corporation','B'),('project','G'),('vendor','I'),('description','S'),('note','V')]} if include_ledger else {})})
    selected=[r for r in ledger if r['department']=='데크거더']
    months=sorted({r['month'] for r in ledger});month_counts=dict(Counter(r['month'] for r in selected))
    adjustments=[]
    for r in [8,9]+EXPENSE_ROWS:
        for col in range(3,15):
            cell=formulas[REPORT].cell(r,col)
            tail=re.search(r'\)((?:[+-]\d+(?:\.\d+)?)+)$',str(cell.value))
            if tail:
                terms=re.findall(r'[+-]\d+(?:\.\d+)?',tail.group(1));amount=sum(map(Decimal,terms),Decimal(0))
                adjustments.append({'cell':cell.coordinate,'row':r,'month':col-2,'amount':float(amount),'formula':cell.value,'terms':terms})
    refunds=[r for r in selected if '환수' in str(r['category'])]
    # Compare the VAT criteria currently used by Excel with a department-scoped variant.
    unsafe_vat=[r for r in ledger if r['department']!='데크거더' and r['category'] in [saved[REPORT].cell(x,2).value for x in EXPENSE_ROWS] and r['vat']!=0]
    for ref in ['C31','O30']:
        amount=number(calc.value(REPORT,ref));cached=number(ws[ref].value)
        checks.append({'cell':ref,'cached':float(cached),'recalculated':float(amount),'difference':float(amount-cached),'ok':abs(amount-cached)<Decimal('0.01')})
    opening=float(number(ws['C30'].value))
    report_by_row={r['row']:r for r in rows}
    breakdowns=[]
    for report_row in [8,9]+EXPENSE_ROWS:
        category='공사수입' if report_row==8 else '제품매출' if report_row==9 else saved[REPORT].cell(report_row,2).value
        for month in months:
            matching=[r for r in selected if r['category']==category and r['month']==month]
            vat_rows=matching if report_row in (8,9) else [r for r in ledger if r['category']==category and r['month']==month]
            breakdowns.append({'report_row':report_row,'month':month,'gross':sum(r['inflow'] if report_row in (8,9) else r['outflow'] for r in matching),'vat':sum(r['vat'] for r in vat_rows),'count':len(matching)})
    classified_outflows=sum(r['outflow'] for r in selected)
    classified_inflows=sum(r['inflow'] for r in selected)
    result={'source':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':REPORT,
            'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in formulas],
            'period_start':min(r['date'] for r in ledger),'period_end':max(r['date'] for r in ledger),'months':months,
            'ledger_count':len(ledger),'selected_count':len(selected),'month_counts':month_counts,
            'rows':rows,'opening':opening,'closing':opening+report_by_row[26]['values'][12],
            'breakdowns':breakdowns,'checks':checks,'adjustments':adjustments,'refunds':refunds,'vat_scope_impact':sum(r['vat'] for r in unsafe_vat),
            'classified_inflows':classified_inflows,'classified_outflows':classified_outflows,
            'refund_total':sum(r['inflow'] for r in refunds),
            'prior_year_claims':ws['B36'].value,
            'notes':[str(ws.cell(r,2).value) for r in [32,33,34] if ws.cell(r,2).value],
            'print_area':str(formulas[REPORT].print_area)}
    if include_ledger:
        result['ledger_rows']=ledger
    return result


def render_html(data, embedded=False):
    payload=json.dumps(data,ensure_ascii=False).replace('</','<\\/')
    template=Path(__file__).with_name('monthly_preview_template.html').read_text(encoding='utf-8')
    result=template.replace('__REPORT_DATA__',payload)
    if embedded:result=result.replace('<body>','<body class="embedded">')
    return result


def choose_file():
    root=Path(__file__).resolve().parent
    candidates=list(root.glob('*통장기준정산*.xlsx'))
    if len(candidates)==1:return candidates[0]
    window=None
    try:
        import tkinter as tk
        from tkinter import filedialog
        window=tk.Tk();window.withdraw();window.attributes('-topmost',True)
        chosen=filedialog.askopenfilename(parent=window,title='통장기준정산 엑셀 선택',filetypes=[('Excel','*.xlsx')])
        if chosen:return Path(chosen)
    except Exception as e:
        print('파일 선택창을 사용할 수 없습니다:',e)
    finally:
        if window:window.destroy()
    value=input('원본 엑셀을 이 창에 끌어다 놓고 Enter를 누르세요: ').strip().strip('"')
    if not value:raise ValueError('파일 선택을 취소했습니다.')
    return Path(value)


def main():
    parser=argparse.ArgumentParser(description='엘비데크 월별보고 분석·시안 생성 (읽기 전용)')
    parser.add_argument('workbook',nargs='?')
    parser.add_argument('--output',type=Path,default=Path(__file__).resolve().parent/'private'/'monthly-preview')
    parser.add_argument('--no-open',action='store_true')
    args=parser.parse_args()
    data=analyze(args.workbook or choose_file())
    args.output.mkdir(parents=True,exist_ok=True)
    (args.output/'analysis.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    output=args.output/'월별보고_분석과시안.html'
    output.write_text(render_html(data),encoding='utf-8')
    print('분석 및 보고서 시안 생성 완료:',output)
    print('원본 대사:',sum(c['ok'] for c in data['checks']),'/',len(data['checks']))
    if not args.no_open:webbrowser.open(output.resolve().as_uri())


if __name__=='__main__':
    try:main()
    except Exception as e:raise SystemExit('분석 실패: '+str(e))
