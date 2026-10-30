"""Read an explicitly mapped worksheet; no workbook instructions are executed."""
from io import BytesIO
from datetime import datetime,date
import openpyxl
from joint_store import normalize

LABELS={'voucher':'전표번호','line':'전표행 번호(연도·번호 조합 권장)','posting_date':'기표일자','allocation_month':'배분월','account':'계정과목','vendor':'거래처','description':'적요','tax_supply':'세금계산서 공급가액','vat':'VAT','invoice':'계산서','receipt':'기타영수증','note':'비고'}
REQUIRED=['voucher','line','posting_date','account','tax_supply']

def read(blob,sheet,header,mapping,end_row=None):
    if any(not mapping.get(k) for k in REQUIRED):raise ValueError('전표번호·전표행·기표일자·계정과목·공급가액 열을 지정하세요.')
    w=openpyxl.load_workbook(BytesIO(blob),data_only=True,read_only=True);f=openpyxl.load_workbook(BytesIO(blob),data_only=False,read_only=True)
    try:
        ws=w[sheet];fs=f[sheet];result=[]
        # Iterate sequentially, avoiding quadratic random access on read-only XML.
        for row,formula_row in zip(ws.iter_rows(min_row=int(header)+1,max_row=end_row),fs.iter_rows(min_row=int(header)+1,max_row=end_row)):
            rec={};used=False
            for key,col in mapping.items():
                if not col:continue
                cell=row[int(col)-1] if int(col)<=len(row) else None
                fc=formula_row[int(col)-1] if int(col)<=len(formula_row) else None
                value=cell.value if cell else None
                if fc is not None and fc.data_type=='f' and value is None:raise ValueError(f'{sheet}!{fc.coordinate}: 수식 계산값이 없습니다. Excel에서 재계산 후 저장하세요.')
                if value is not None:used=True
                if key=='posting_date' and value is not None:
                    if isinstance(value,(date,datetime)):value=value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
                    else:
                        value=str(value).strip().replace('.','-').replace('/','-')
                        value=datetime.strptime(value,'%Y-%m-%d').date().isoformat()
                if key=='allocation_month' and isinstance(value,(date,datetime)):value=value.strftime('%Y-%m')
                rec[key]='' if value is None else str(value)
            if not used:continue
            try:result.append(dict(normalize(rec),voucher=rec.get('voucher',''),line=rec.get('line','')))
            except ValueError as e:raise ValueError(f'{row[0].row}행: {e}')
        return result
    finally:w.close();f.close()
